"""Script to populate the SLA & Settings tab in Member.xlsx."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pathlib import Path

SETTINGS_DATA = [
    # Category, Setting Name, Setting Key, Value, Unit/Options, Description
    ("SLA Timeline", "Stage 1 TL Follow-up (Minutes)", "STAGE1_TL_MINUTES", 10, "Minutes (e.g. 10)", "Minutes of silence before tagging Team Lead on Teams"),
    ("SLA Timeline", "Enable Stage 1 Follow-up", "STAGE1_ENABLED", True, "TRUE / FALSE", "Enable or disable Stage 1 Team Lead follow-up"),
    ("SLA Timeline", "Stage 2 PM Escalation (Minutes)", "STAGE2_PM_MINUTES", 20, "Minutes (e.g. 20)", "Minutes of silence before escalating to Project Manager on Teams"),
    ("SLA Timeline", "Enable Stage 2 Escalation", "STAGE2_ENABLED", True, "TRUE / FALSE", "Enable or disable Stage 2 PM escalation card"),
    ("SLA Timeline", "Stage 3 HM Email Alert (Minutes)", "STAGE3_HM_EMAIL_MINUTES", 20, "Minutes (e.g. 20)", "Minutes of silence before sending urgent alert email to Higher Management"),
    ("SLA Timeline", "Enable Stage 3 Email Alert", "STAGE3_ENABLED", True, "TRUE / FALSE", "Enable or disable Stage 3 Higher Management email alerts"),
    ("SLA Timeline", "Ignore Ancient Backlog (Hours)", "BACKLOG_CUTOFF_HOURS", 24, "Hours (e.g. 24)", "Suppress reminders for unaddressed messages older than X hours"),
    ("SLA Timeline", "SLA Monitor Check Interval (Sec)", "REMINDER_CHECK_INTERVAL", 30, "Seconds (e.g. 30)", "How frequently the background engine checks message SLAs"),

    ("Business Hours", "Business Hours Only", "BUSINESS_HOURS_ONLY", False, "TRUE / FALSE", "If TRUE, reminders are only dispatched during working hours"),
    ("Business Hours", "Workday Start Time", "BUSINESS_START_TIME", "09:00", "HH:MM (24-hour)", "Start of business hours (e.g. 09:00 AM)"),
    ("Business Hours", "Workday End Time", "BUSINESS_END_TIME", "19:00", "HH:MM (24-hour)", "End of business hours (e.g. 07:00 PM)"),
    ("Business Hours", "Exclude Weekends", "EXCLUDE_WEEKENDS", True, "TRUE / FALSE", "Pause follow-up reminders on Saturday and Sunday"),
    ("Business Hours", "Timezone", "TIMEZONE", "Asia/Kolkata", "Timezone name", "Timezone for business hours evaluation"),

    ("Jira Defaults", "Default Jira Project Key", "DEFAULT_JIRA_PROJECT", "SCRUM", "Project Key", "Fallback Jira Project Key for ticket creation"),
    ("Jira Defaults", "Default Issue Type", "DEFAULT_ISSUE_TYPE", "Bug", "Bug / Task / Story", "Default Jira issue type"),
    ("Jira Defaults", "Default Priority", "DEFAULT_PRIORITY", "Medium", "Lowest / Low / Medium / High / Highest", "Default Jira ticket priority"),
    ("Jira Defaults", "Auto-Approve on TL/PM Reaction", "AUTO_APPROVE_FOR_TL_PM", False, "TRUE / FALSE", "If TRUE, reaction directly creates ticket without confirmation card"),
    ("Jira Defaults", "Approval Reactions", "REQUIRE_APPROVAL_REACTIONS", "🎟️, 🎫, 👍, like", "Comma-separated", "Reaction emojis that authorize ticket confirmation"),
    ("Jira Defaults", "Disapproval Reactions", "DISAPPROVAL_REACTIONS", "❌, decline, thumbs_down", "Comma-separated", "Reaction emojis that reject ticket creation"),

    ("AI & Triggers", "Issue Trigger Keywords", "TRIGGER_KEYWORDS", "#issue, #bug, #task, bug, issue, error, fail, broken, crash, not working, 500 error", "Comma-separated", "Keywords that cause incoming client messages to be treated as bugs/tasks"),
    ("AI & Triggers", "Ignored Keywords", "IGNORE_KEYWORDS", "fyi, resolved, fixed, false alarm, testing, disregard", "Comma-separated", "Keywords that suppress issue creation"),
]

def update_member_workbook(p: Path):
    if not p.exists():
        return
    wb = openpyxl.load_workbook(p)
    if "SLA & Settings" in wb.sheetnames:
        del wb["SLA & Settings"]
    elif "Sheet1" in wb.sheetnames:
        del wb["Sheet1"]
    ws = wb.create_sheet("SLA & Settings")

    # Title Banner
    ws.merge_cells("A1:F1")
    ws["A1"] = "Jira Automation — System SLA Timelines, Business Rules & Operational Policies"
    ws["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:F2")
    ws["A2"] = "Single Source of Truth for SLA Timing, Escalation Hierarchy, Working Hours & Jira Defaults (Auto-synced)"
    ws["A2"].font = Font(name="Segoe UI", size=9.5, italic=True, color="94A3B8")
    ws["A2"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Header Row
    headers = ["Category", "Setting Name", "Setting Key", "Value", "Unit / Options", "Description & Guidance"]
    ws.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.alignment = Alignment(horizontal="left" if col_idx != 4 else "center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    dv_bool = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
    ws.add_data_validation(dv_bool)

    cat_colors = {
        "SLA Timeline": ("EEF2FF", "312E81"),
        "Business Hours": ("ECFDF5", "065F46"),
        "Jira Defaults": ("F0FDF4", "166534"),
        "AI & Triggers": ("FAF5FF", "6B21A8"),
    }

    for idx, (cat, name, key, val, unit, desc) in enumerate(SETTINGS_DATA, start=4):
        ws.row_dimensions[idx].height = 24
        is_even = idx % 2 == 0
        bg_color = "F8FAFC" if is_even else "FFFFFF"

        ws.cell(row=idx, column=1, value=cat)
        ws.cell(row=idx, column=2, value=name)
        ws.cell(row=idx, column=3, value=key)

        val_cell = ws.cell(row=idx, column=4)
        if isinstance(val, bool):
            val_cell.value = "TRUE" if val else "FALSE"
            dv_bool.add(val_cell)
        else:
            val_cell.value = val

        ws.cell(row=idx, column=5, value=unit)
        ws.cell(row=idx, column=6, value=desc)

        cat_bg, cat_fg = cat_colors.get(cat, ("F1F5F9", "334155"))

        for c_idx in range(1, 7):
            c = ws.cell(row=idx, column=c_idx)
            c.border = thin_border
            if c_idx == 1:
                c.fill = PatternFill(start_color=cat_bg, end_color=cat_bg, fill_type="solid")
                c.font = Font(name="Segoe UI", size=9.5, bold=True, color=cat_fg)
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                c.font = Font(name="Consolas", size=9.5, bold=True, color="0F172A")
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 4:
                is_num = "MINUTES" in key or "INTERVAL" in key or "HOURS" in key
                c.fill = PatternFill(start_color="FEF3C7" if is_num else ("DCFCE7" if str(val) in ("TRUE", "SCRUM", "Bug") else "F1F5F9"), fill_type="solid")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
                c.alignment = Alignment(horizontal="center" if len(str(val)) < 15 else "left", vertical="center")
            else:
                c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                c.font = Font(name="Segoe UI", size=9.5, color="334155")
                c.alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 24
    ws.column_dimensions["F"].width = 65

    wb.save(p)
    print("Successfully updated SLA & Settings tab in:", str(p))

if __name__ == "__main__":
    paths = [
        Path.home() / "OneDrive" / "Jira-Automation" / "Member.xlsx",
        Path("data/Member.xlsx"),
    ]
    for p in paths:
        update_member_workbook(p)
